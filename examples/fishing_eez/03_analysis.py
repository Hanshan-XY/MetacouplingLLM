#!/usr/bin/env python
"""
03_analysis.py  --  Reproduce the paper's regional tables (Tables 4 and 5) and the
summary figures of its fishing example (Section 5.2).

INPUT
    outputs/indicators_by_eez.csv  -- per-EEZ indicators from the package (step 02)
    data/fishing_edge_table.csv    -- the edge table (for nominal partner counts)
    data/eez_region.json           -- per-EEZ region map  {EEZ: {grp, region}}

OUTPUT
    outputs/regional_catch.csv       -- Table 4: regional catch totals and proportions
    outputs/regional_indicators.csv  -- Table 5: catch-weighted regional indicators
    outputs/global_totals.csv        -- global shares and the Section 5.2 figures
    outputs/telecoupled_share_timeseries.csv   (optional; needs the source tables)

ALL the statistics in this file are USER-SUPPLIED analysis built ON TOP of the
package's per-EEZ indicator table.  The package computes the indicators; this
script aggregates them by region.

  * Regions follow the 'grp' field of eez_region.json (Carlson et al.'s Table S16
    scheme, joined by 01_build_region_map.py).  The eleven regions cover 276 of the
    280 EEZs; four EEZs absent from Carlson's nation list are unassigned (Taiwan,
    Cape Verde, Cook Islands, Ascension) and are left out of both tables.
  * Regional indicator means are CATCH-WEIGHTED (weight = F_total).  An EEZ with no
    flow in a coupling type has no PFCI/TFCI/ENP for it, so those means cover only
    EEZs with positive catch in that type, with the weights renormalized.
  * Rows are ordered by telecoupled share, highest first, as in the paper.
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
IND = HERE / "outputs" / "indicators_by_eez.csv"
EDGES = HERE / "data" / "fishing_edge_table.csv"
REGION_JSON = HERE / "data" / "eez_region.json"
OUT_DIR = HERE / "outputs"
SRC_DIR = Path(os.environ.get("CARLSON2020_SUPPLEMENT", HERE / "source"))

# Columns reported in Table 5 (the catch-weighted regional means).
INDICATOR_COLS = ["IFS", "PFS", "TFS", "MFE", "PFCI", "TFCI", "ENP_P", "ENP_T"]


def weighted_mean(values: pd.Series, weights: pd.Series) -> float:
    """Catch-weighted mean, ignoring NaN indicator values (and their weight)."""
    v = values.to_numpy(dtype=float)
    w = weights.to_numpy(dtype=float)
    mask = ~np.isnan(v)
    if not mask.any() or w[mask].sum() == 0:
        return np.nan
    return float(np.average(v[mask], weights=w[mask]))


def main() -> None:
    ind = pd.read_csv(IND)
    edges = pd.read_csv(EDGES)

    region_map = json.loads(REGION_JSON.read_text(encoding="utf-8"))
    grp = {eez: meta.get("grp") for eez, meta in region_map.items()}
    ind["region"] = ind["focal_system_id"].map(grp)
    reg = ind[ind["region"].notna()].copy()

    # ----- Table 4: regional catch totals and proportions (million metric tons) ---
    catch_rows = []
    for region, sub in reg.groupby("region"):
        total, t1, t2, t3 = (sub[c].sum() / 1e6 for c in ("F_total", "F_I", "F_P", "F_T"))
        catch_rows.append({
            "region": region, "n_eez": len(sub), "total_MMT": total,
            "type1_MMT": t1, "type2_MMT": t2, "type3_MMT": t3,
            "type1_pct": 100 * t1 / total, "type2_pct": 100 * t2 / total,
            "type3_pct": 100 * t3 / total,
        })
    table4 = pd.DataFrame(catch_rows).sort_values("type3_pct", ascending=False)
    table4.to_csv(OUT_DIR / "regional_catch.csv", index=False)

    # ----- Table 5: catch-weighted regional indicators ---------------------------
    ind_rows = []
    for region, sub in reg.groupby("region"):
        row = {"region": region, "n_eez": len(sub)}
        for col in INDICATOR_COLS:
            row[col] = weighted_mean(sub[col], sub["F_total"])
        ind_rows.append(row)
    table5 = pd.DataFrame(ind_rows).sort_values("TFS", ascending=False)
    table5.to_csv(OUT_DIR / "regional_indicators.csv", index=False)

    # ----- Global figures --------------------------------------------------------
    total_t = ind["F_total"].sum()
    distant = ind[ind["F_T"] > 0]
    nominal = (edges[edges["coupling_type"] == "T"]
               .groupby("focal_system_id")["destination_id"].nunique()
               .reindex(distant["focal_system_id"]))
    global_rows = pd.DataFrame([
        {"metric": "n_eez", "value": ind["focal_system_id"].nunique()},
        {"metric": "total_catch_Bt", "value": total_t / 1e9},
        {"metric": "own_plus_adjacent_share", "value": (ind["F_I"].sum() + ind["F_P"].sum()) / total_t},
        {"metric": "telecoupled_share", "value": ind["F_T"].sum() / total_t},
        {"metric": "n_eez_region_assigned", "value": len(reg)},
        {"metric": "region_assigned_catch_share", "value": reg["F_total"].sum() / total_t},
        {"metric": "n_eez_with_distant_catch", "value": len(distant)},
        {"metric": "mean_ENP_T_unweighted", "value": distant["ENP_T"].mean()},
        {"metric": "mean_nominal_distant_nations", "value": nominal.mean()},
    ])
    global_rows.to_csv(OUT_DIR / "global_totals.csv", index=False)

    g = dict(zip(global_rows["metric"], global_rows["value"]))
    print("=== Global ===")
    print(f"  EEZs: {g['n_eez']:.0f}; region-assigned: {g['n_eez_region_assigned']:.0f} "
          f"({100 * g['region_assigned_catch_share']:.2f}% of catch)")
    print(f"  total catch {g['total_catch_Bt']:.3f} Bt; own + adjacent {100 * g['own_plus_adjacent_share']:.1f}%")
    print(f"  EEZs with distant catch: {g['n_eez_with_distant_catch']:.0f}; unweighted mean ENP_T "
          f"{g['mean_ENP_T_unweighted']:.2f} vs {g['mean_nominal_distant_nations']:.2f} nominal distant nations")
    with pd.option_context("display.width", 200, "display.max_columns", 20):
        print("\n=== Table 4: regional catch (MMT, %) ===")
        print(table4.set_index("region").round(1))
        print("\n=== Table 5: catch-weighted regional indicators ===")
        print(table5.set_index("region").round(2))

    _telecoupled_timeseries()


def _telecoupled_timeseries() -> None:
    """Per-year global telecoupled share = Type3 / (Type1+Type2+Type3).

    Recomputes annual catch totals straight from the source tables (the pinned
    edge table sums over years, so it cannot give a per-year series).  This is a
    supplementary, descriptive figure -- not part of the regional tables.
    """
    import openpyxl

    if not (SRC_DIR / "Table S4.xlsx").exists():
        print(f"\n(Source tables not found in {SRC_DIR}; skipping temporal series.)")
        return

    def s4_year_totals(path):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["catch"]
        totals = {}
        for row in ws.iter_rows(min_row=4, values_only=True):
            year = row[0]
            if year is None:
                continue
            s = sum(float(v) for v in row[1:] if v is not None)
            totals[int(float(year))] = totals.get(int(float(year)), 0.0) + s
        wb.close()
        return totals

    def block_year_totals(path):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["catch"]
        header = next(ws.iter_rows(min_row=3, max_row=3, values_only=True))
        # catch lives in columns whose header == 'Catch'
        catch_cols = [c for c, h in enumerate(header) if h is not None and str(h).strip().lower() == "catch"]
        totals = {}
        for row in ws.iter_rows(min_row=4, values_only=True):
            year = row[0]
            if year is None:
                continue
            s = sum(float(row[c]) for c in catch_cols if c < len(row) and row[c] is not None)
            totals[int(float(year))] = totals.get(int(float(year)), 0.0) + s
        wb.close()
        return totals

    t1 = s4_year_totals(SRC_DIR / "Table S4.xlsx")
    t2 = block_year_totals(SRC_DIR / "Table S6.xlsx")
    t3 = block_year_totals(SRC_DIR / "Table S8.xlsx")

    years = sorted(set(t1) | set(t2) | set(t3))
    rows = []
    for y in years:
        a, b, c = t1.get(y, 0.0), t2.get(y, 0.0), t3.get(y, 0.0)
        tot = a + b + c
        rows.append({"year": y, "telecoupled_share": (c / tot) if tot else np.nan})
    ts = pd.DataFrame(rows)
    ts.to_csv(OUT_DIR / "telecoupled_share_timeseries.csv", index=False)

    peak = ts.loc[ts["telecoupled_share"].idxmax()]
    last = ts.iloc[-1]
    print("\n=== Telecoupled share over time (optional) ===")
    print(f"  peak: {peak['telecoupled_share']:.3f} in {int(peak['year'])}")
    print(f"  {int(last['year'])}: {last['telecoupled_share']:.3f}")


if __name__ == "__main__":
    main()
