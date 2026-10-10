#!/usr/bin/env python
"""
00_build_edge_table.py  --  Build the metacoupling edge table for global marine
fishing across EEZs (1950-2014).

INPUT  (supplementary tables of Carlson et al., 2020; Sea Around Us reconstructions),
read from the folder named by the CARLSON2020_SUPPLEMENT environment variable, or
from ./source when it is unset:
    Table S4.xlsx  -- Type 1 catch: an EEZ fishing its OWN waters        (intracoupling, I)
    Table S6.xlsx  -- Type 2 catch: ADJACENT nations fishing the EEZ     (pericoupling,  P)
    Table S8.xlsx  -- Type 3 catch: DISTANT  nations fishing the EEZ     (telecoupling,  T)

    Each .xlsx has a 'catch' sheet.  Row 3 is the header; data start at row 4.
    Column A holds the year (1950-2014).

      * S4 is WIDE  : one column per EEZ, the cell value is that EEZ's own catch.
      * S6 / S8 are BLOCKS : the header is a repeating pair (EEZ-name, 'Catch').
        For each EEZ the left column holds the *fishing-nation* name and the right
        column the catch.  Within a block the same year repeats once per fishing
        nation active that year, so we aggregate (sum) over all years AND over
        repeated nation rows.

OUTPUT (pinned so downstream steps run offline):
    data/fishing_edge_table.csv  -- one row per flow, columns:
        focal_system_id, destination_id, coupling_type {I,P,T}, flow_value

    Edge construction:
        Type 1 -> {focal: EEZ, destination_id: EEZ+' (own)', coupling_type:'I', flow_value: own catch}
        Type 2 -> one row per adjacent nation {focal: EEZ, destination_id: nation, 'P', catch}
        Type 3 -> one row per distant  nation {focal: EEZ, destination_id: nation, 'T', catch}

This script does NOT use any metacouplingllm functionality -- it is pure
data-wrangling (USER-SUPPLIED) that produces the package's expected input format.
"""

from __future__ import annotations

import csv
import os
from collections import defaultdict
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
SRC_DIR = Path(os.environ.get("CARLSON2020_SUPPLEMENT", HERE / "source"))
DATA_DIR = HERE / "data"
DATA_DIR.mkdir(exist_ok=True)

# Two EEZs carry one name in the own-catch table (S4) and another in the
# foreign-catch tables (S6/S8); unmerged, each would split into an all-domestic
# and an all-foreign focal system.
EEZ_ALIASES = {
    "Aruba (Netherlands)": "Aruba",
    "Congo, R. of": "Congo (Republic of)",
}

S4 = SRC_DIR / "Table S4.xlsx"   # Type 1 (own / intracoupling)
S6 = SRC_DIR / "Table S6.xlsx"   # Type 2 (adjacent / pericoupling)
S8 = SRC_DIR / "Table S8.xlsx"   # Type 3 (distant  / telecoupling)

HEADER_ROW = 3        # row holding EEZ / 'Catch' column labels
FIRST_DATA_ROW = 4    # first year row


def _clean(name) -> str:
    """Trim whitespace; several EEZ/nation labels carry trailing spaces."""
    return str(name).strip() if name is not None else ""


def _eez(name) -> str:
    eez = _clean(name)
    return EEZ_ALIASES.get(eez, eez)


def parse_type1_own(path: Path) -> dict[str, float]:
    """S4 (wide): sum each EEZ column over all year rows -> own catch per EEZ."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["catch"]

    header = next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))
    # column 0 is 'Year/EEZ'; columns 1.. are EEZ names
    eez_by_col = {col: _eez(name) for col, name in enumerate(header) if col >= 1 and name}

    totals: dict[str, float] = defaultdict(float)
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        for col, eez in eez_by_col.items():
            val = row[col]
            if val is not None:
                totals[eez] += float(val)
    wb.close()
    return dict(totals)


def parse_blocks(path: Path) -> dict[str, dict[str, float]]:
    """S6 / S8 (blocks): for each EEZ, sum catch per fishing-nation over all years.

    Returns {EEZ: {fishing_nation: total_catch, ...}, ...}.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["catch"]

    header = next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))
    # Header is a repeating pair: (EEZ-name col, 'Catch' col).  The EEZ name sits
    # in even-offset columns starting at index 1; its catch is the next column.
    eez_block = {}  # name_col_index -> EEZ name
    col = 1
    n = len(header)
    while col < n:
        label = _clean(header[col])
        if label and label.lower() != "catch":
            eez_block[col] = EEZ_ALIASES.get(label, label)
        col += 1
    # Keep only the columns that are EEZ-name columns (those whose +1 neighbour is 'Catch')
    name_cols = [c for c in eez_block if (c + 1) < n and _clean(header[c + 1]).lower() == "catch"]

    result: dict[str, dict[str, float]] = {eez_block[c]: defaultdict(float) for c in name_cols}
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        for c in name_cols:
            nation = row[c]
            catch = row[c + 1] if (c + 1) < len(row) else None
            if nation is not None and catch is not None:
                eez = eez_block[c]
                result[eez][_clean(nation)] += float(catch)

    wb.close()
    # drop empties, demote defaultdicts to plain dicts
    return {eez: dict(d) for eez, d in result.items() if d}


def main() -> None:
    missing = [p.name for p in (S4, S6, S8) if not p.exists()]
    if missing:
        raise SystemExit(
            f"Source tables not found in {SRC_DIR}: {', '.join(missing)}. Put them in "
            "./source or set CARLSON2020_SUPPLEMENT (see README); the pinned "
            "data/fishing_edge_table.csv works without them."
        )
    print("Parsing Type 1 (own / I) from", S4.name)
    own = parse_type1_own(S4)
    print(f"  {len(own)} EEZs with own catch")

    print("Parsing Type 2 (adjacent / P) from", S6.name)
    peri = parse_blocks(S6)
    print(f"  {len(peri)} EEZs with adjacent-nation catch")

    print("Parsing Type 3 (distant / T) from", S8.name)
    tele = parse_blocks(S8)
    print(f"  {len(tele)} EEZs with distant-nation catch")

    # Universe of EEZs is the union across all three tables.
    all_eez = sorted(set(own) | set(peri) | set(tele))

    edges: list[dict] = []
    for eez in all_eez:
        # Type 1 -> single self-loop edge (own waters).
        if eez in own and own[eez] != 0:
            edges.append(
                {
                    "focal_system_id": eez,
                    "destination_id": f"{eez} (own)",
                    "coupling_type": "I",
                    "flow_value": own[eez],
                }
            )
        # Type 2 -> one edge per adjacent fishing nation.
        for nation, catch in peri.get(eez, {}).items():
            if catch != 0:
                edges.append(
                    {
                        "focal_system_id": eez,
                        "destination_id": nation,
                        "coupling_type": "P",
                        "flow_value": catch,
                    }
                )
        # Type 3 -> one edge per distant fishing nation.
        for nation, catch in tele.get(eez, {}).items():
            if catch != 0:
                edges.append(
                    {
                        "focal_system_id": eez,
                        "destination_id": nation,
                        "coupling_type": "T",
                        "flow_value": catch,
                    }
                )

    out = DATA_DIR / "fishing_edge_table.csv"
    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f, fieldnames=["focal_system_id", "destination_id", "coupling_type", "flow_value"]
        )
        w.writeheader()
        w.writerows(edges)

    n_eez = len({e["focal_system_id"] for e in edges})
    print(f"\nWrote {len(edges)} edge rows across {n_eez} EEZs -> {out}")


if __name__ == "__main__":
    main()
